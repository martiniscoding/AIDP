"""Cells in, lines and tables out — the adapter for Excel workbooks.

A workbook is the odd one out in this corpus. A PDF is prose whose structure was
destroyed at export and has to be rebuilt from typography; a deck is prose in
boxes. A spreadsheet is not prose at all — most of what it says, it says in a
grid, and the grid *is* the meaning. A criticality table whose Tier 4 row leaves
RPO and RTO blank says something precise, and flattening it to a sentence turns
"not specified" into the previous tier's value.

So this splits each sheet in two rather than turning it all into text:

  * blocks that look like a table    -> `tables.Table`, header bound to every
                                        cell, empty cells kept as None. The
                                        chunker already emits one chunk per row
                                        with the headers inline, which is
                                        exactly what a row needs to be
                                        retrievable on its own.
  * everything else                  -> lines: the sheet's title rows, notes
                                        under a table, a stray cell of prose.

Table rows are deliberately *not* also emitted as lines. They are already
chunked one per row, and duplicating them into the section's prose would put
the same sentence in the index twice and let a single fact answer as two
independent pieces of evidence.

What this does not do is decide the document's structure. A sheet name is
frequently not a clause heading — a standards workbook says "Ctrl-14" where a
standard says "All data at rest must be encrypted" — so the same model pass
that rescues a formatless PDF (`ai_structure`) decides here too, exactly as it
does for a deck. This module's job is to hand that pass the cleanest possible
input, with each line tagged with what the workbook already knew about it.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook

from .spans import Line
from .tables import Table

# A guard, not a judgement about how big a real workbook gets. An export with a
# hundred thousand rows is a data dump rather than a document, and reading it
# whole would cost more memory than the worker has. What is read is reported.
MAX_ROWS_PER_SHEET = 5_000
MAX_COLUMNS = 64
# Below this a block is prose that happens to sit in adjacent cells, not a
# table: one column is a list, and one row is a heading.
MIN_TABLE_ROWS = 2
MIN_TABLE_COLUMNS = 2

_SIZE_SHEET_NAME = 18.0
_SIZE_BODY = 11.0


@dataclass
class Workbook:
    """Shaped to drop straight into the parse stage, like `slides.Deck`."""

    lines: list[Line] = field(default_factory=list)
    #: Line index -> what the workbook already knew this row was. Passed to the
    #: structure pass as a hint; see `_tagged` in ai_structure.py.
    hints: dict[int, str] = field(default_factory=dict)
    #: Keyed the way `Figure.page` is on the PDF path — here the sheet number.
    tables: list[Table] = field(default_factory=list)
    sheet_count: int = 0
    #: The workbook's own name for itself, when the first sheet opens with one.
    title: str | None = None
    #: Sheets whose rows ran past MAX_ROWS_PER_SHEET, reported rather than
    #: silently truncated.
    truncated: list[str] = field(default_factory=list)


def _sig(size: float, bold: bool) -> tuple[str, float, bool, bool, int]:
    return ("sheet", size, bold, False, 0)


def _clean(value: Any) -> str | None:
    """A cell as text, or None when it holds nothing.

    None and empty are the same thing to a reader and must stay distinct from
    "0" or "False", which are values a spreadsheet means.
    """
    if value is None:
        return None
    if isinstance(value, str):
        text = " ".join(value.split())
        return text or None
    return str(value)


def _emit(
    book: Workbook, text: str, *, sheet: int, hint: str, size: float, bold: bool, y: float
) -> None:
    body = " ".join(text.split())
    if not body:
        return
    index = len(book.lines)
    book.lines.append(
        Line(
            text=body,
            page=sheet,
            y=y,
            x0=0.0,
            bbox=(0.0, y, 0.0, y),
            sig=_sig(size, bold),
            has_bold=bold,
        )
    )
    book.hints[index] = hint


def _blocks(rows: list[list[str | None]]) -> list[tuple[int, list[list[str | None]]]]:
    """Consecutive non-empty rows, grouped, with the offset each group started at.

    A blank row is how a spreadsheet says "new thing" — it is the only
    separator the format has, and it is used consistently enough to be worth
    trusting.
    """
    out: list[tuple[int, list[list[str | None]]]] = []
    current: list[list[str | None]] = []
    start = 0
    for index, row in enumerate(rows):
        if any(cell is not None for cell in row):
            if not current:
                start = index
            current.append(row)
        elif current:
            out.append((start, current))
            current = []
    if current:
        out.append((start, current))
    return out


def _width(block: list[list[str | None]]) -> int:
    """How many columns the block actually occupies."""
    widest = 0
    for row in block:
        for column, cell in enumerate(row):
            if cell is not None:
                widest = max(widest, column + 1)
    return widest


def _columns(header: list[str | None], width: int) -> list[str]:
    """Column names, with a placeholder wherever the sheet left one blank.

    A nameless column still holds values, and dropping it would shift every
    cell after it — the failure this whole module exists to avoid.
    """
    names: list[str] = []
    seen: dict[str, int] = {}
    for index in range(width):
        raw = header[index] if index < len(header) else None
        name = raw or f"Column {index + 1}"
        # Two columns called "Notes" would collapse into one key and the second
        # would overwrite the first.
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        names.append(name)
    return names


def _looks_tabular(block: list[list[str | None]], width: int) -> bool:
    if len(block) < MIN_TABLE_ROWS or width < MIN_TABLE_COLUMNS:
        return False
    # A header row worth the name fills most of its columns. Without that the
    # block is a paragraph split across cells, and binding the rest of it to
    # whatever happened to sit in row one would invent a meaning.
    header = block[0]
    filled = sum(1 for index in range(width) if index < len(header) and header[index] is not None)
    return filled >= max(MIN_TABLE_COLUMNS, (width + 1) // 2)


def _read_sheet(book: Workbook, worksheet, *, sheet: int) -> None:
    rows: list[list[str | None]] = []
    for raw in worksheet.iter_rows(values_only=True):
        if len(rows) >= MAX_ROWS_PER_SHEET:
            book.truncated.append(worksheet.title)
            break
        rows.append([_clean(cell) for cell in raw[:MAX_COLUMNS]])

    # The sheet's name is the one piece of structure the format states outright.
    _emit(book, worksheet.title, sheet=sheet, hint="sheet", size=_SIZE_SHEET_NAME, bold=True, y=0.0)

    for offset, block in _blocks(rows):
        width = _width(block)
        if _looks_tabular(block, width):
            columns = _columns(block[0], width)
            body: list[dict[str, str | None]] = []
            for row in block[1:]:
                padded = list(row) + [None] * (width - len(row))
                body.append({columns[i]: padded[i] for i in range(width)})
            if body:
                book.tables.append(
                    Table(
                        page_start=sheet,
                        page_end=sheet,
                        bbox=(0.0, float(offset), float(width), float(offset + len(block))),
                        columns=columns,
                        rows=body,
                        confidence=1.0,
                        caption=worksheet.title,
                    )
                )
                continue

        # Not a table: prose that happens to live in cells. Each row becomes one
        # line, its cells joined in reading order.
        for index, row in enumerate(block):
            text = " ".join(cell for cell in row if cell)
            _emit(
                book,
                text,
                sheet=sheet,
                hint="text",
                size=_SIZE_BODY,
                bold=False,
                y=float(offset + index + 1),
            )


def read(raw: bytes) -> Workbook:
    """Parse a .xlsx into lines and tables.

    `data_only` so a formula yields the value Excel last computed rather than
    "=SUM(B2:B9)", which is not what the sheet says to a reader. A workbook
    saved by a tool that never calculated will have None there, and those cells
    read as empty — which is honest: we do not evaluate formulas.
    """
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    book = Workbook()
    try:
        sheets = workbook.worksheets
        book.sheet_count = len(sheets)
        for number, worksheet in enumerate(sheets, start=1):
            _read_sheet(book, worksheet, sheet=number)
    finally:
        workbook.close()

    # The first written line of the first sheet, when it is not simply the
    # sheet's own name. Indexed by position rather than searched for by value:
    # two sheets can hold the same text, and `list.index` would find the wrong
    # one of them.
    for index, line in enumerate(book.lines):
        if line.page != 1:
            break
        if book.hints.get(index) == "text":
            book.title = line.text[:200]
            break

    return book

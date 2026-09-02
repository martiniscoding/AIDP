"""Smoke test for the workbook path — no database, no network.

Builds a spreadsheet carrying the things that actually go wrong in one, then
runs the adapter over it and asserts on the result:

  · a title row, a blank, then a grid    → the title stays prose, the grid
                                           becomes a table, and the blank row
                                           is what separates them
  · a row with two empty cells mid-grid  → those cells stay empty rather than
                                           every value after them sliding left.
                                           This is the whole reason the adapter
                                           does not flatten a sheet to text: a
                                           Tier 4 row with no RPO is a fact, and
                                           "4 to 8 hours" borrowed from Tier 3
                                           is a wrong recovery objective
                                           delivered with a citation
  · two columns with the same name       → kept apart, not collapsed onto one
                                           key where the second overwrites the
                                           first
  · a note under the grid                → kept as prose
  · table rows                           → NOT also emitted as lines, or the
                                           same fact would be indexed twice and
                                           could answer as two independent
                                           pieces of evidence
  · a second sheet                       → its table anchored to sheet 2, the
                                           way a figure is anchored to its page

Run inside the worker image, which already has openpyxl:

    python -m scripts.smoke_sheets
"""

from __future__ import annotations

import io
import sys

from openpyxl import Workbook as XlWorkbook

from aidp.parsing import sheets

failures: list[str] = []


def ok(name: str, condition: bool, extra: object = "") -> None:
    print(f"  {'PASS' if condition else 'FAIL'} {name}" + ("" if condition else f"  {extra!r}"))
    if not condition:
        failures.append(name)


def build() -> bytes:
    wb = XlWorkbook()
    ws = wb.active
    ws.title = "Criticality"
    ws.append(["Data Standards — Criticality Tiers"])
    ws.append([])
    ws.append(["Tier", "Description", "RPO", "RTO", "Owner", "Notes", "Notes"])
    ws.append(
        ["Tier 1", "Mission critical", "0 minutes", "15 minutes", "Platform", "a", "b"]
    )
    ws.append(
        ["Tier 2", "Business critical", "15 minutes", "1 hour", "Platform", None, None]
    )
    ws.append(["Tier 3", "Important", "4 hours", "8 hours", "Apps", None, None])
    ws.append(["Tier 4", "Low", None, None, "Apps", None, None])
    ws.append([])
    ws.append(["Note: tiers are reviewed annually by the ARB."])

    controls = wb.create_sheet("Controls")
    controls.append(["Ctrl", "Requirement"])
    controls.append(
        ["C-01", "All data at rest must be encrypted with organisation-managed keys."]
    )
    controls.append(["C-02", "MFA is required for administrative access."])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


book = sheets.read(build())

print("\n1. Shape")
ok("two sheets", book.sheet_count == 2, book.sheet_count)
ok("one table per sheet", len(book.tables) == 2, len(book.tables))
ok("title read off the title row", book.title == "Data Standards — Criticality Tiers", book.title)

print("\n2. The grid keeps its meaning")
table = book.tables[0]
ok(
    "headers bound to columns",
    table.columns[:5] == ["Tier", "Description", "RPO", "RTO", "Owner"],
    table.columns,
)
ok("duplicate header kept apart", table.columns[5] != table.columns[6], table.columns[5:])
ok("four data rows", len(table.rows) == 4, len(table.rows))

tier4 = table.rows[3]
ok("Tier 4 is still the fourth row", tier4["Tier"] == "Tier 4", tier4)
ok("its empty RPO stays empty", tier4["RPO"] is None, tier4["RPO"])
ok("its empty RTO stays empty", tier4["RTO"] is None, tier4["RTO"])
ok("the value after them did not slide left", tier4["Owner"] == "Apps", tier4["Owner"])
ok(
    "a row reads as 'not specified'",
    "RPO: not specified" in table.render_row(tier4),
    table.render_row(tier4),
)

print("\n3. Prose is kept, and kept out of the grid")
texts = [line.text for line in book.lines]
ok("sheet names became lines", "Criticality" in texts and "Controls" in texts, texts[:3])
ok(
    "the note under the grid survived",
    any(t.startswith("Note: tiers") for t in texts),
    texts,
)
ok(
    "table rows are not duplicated into prose",
    not any("Mission critical" in t for t in texts),
    [t for t in texts if "Mission" in t],
)

print("\n4. The structure pass gets the format's own tags")
hints = {book.hints[i] for i in range(len(book.lines))}
ok("sheet names hinted", "sheet" in hints, hints)
ok("prose hinted", "text" in hints, hints)

print("\n5. Sheets anchor their tables the way pages anchor figures")
ok("second table sits on sheet 2", book.tables[1].page_start == 2, book.tables[1].page_start)
ok("and kept both its rows", len(book.tables[1].rows) == 2, book.tables[1].rows)

print(f"\n{len(failures)} failed" if failures else "\nall passed")
sys.exit(1 if failures else 0)

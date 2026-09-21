"""Slides drawn as diagrams, photographed — and only those.

A deck is built here with python-pptx, one slide of each kind that matters: plain
bullets, a diagram of boxes joined by connectors, text laid out in boxes with no
arrows, a pasted picture, and a hidden slide in front of the diagram so that a
PDF missing it would shift every page. What this proves: only the drawn diagram
is picked; its picture is taken of the right slide, through LibreOffice, and
added as a figure with the slide's title; and when LibreOffice is missing, stuck
or disagrees about the page count, the deck is still read and says why its
diagram has no picture.

Needs LibreOffice for the rendering half, so it runs in the worker image:

    docker run --rm -v "$PWD/Workers":/w -w /w -e STAGE=parse \\
      -e DATABASE_URL=postgresql://nobody@127.0.0.1:1/none \\
      --entrypoint python workers-parse scripts/smoke_slide_diagrams.py
"""

from __future__ import annotations

import io
import pathlib
import subprocess
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

import fitz  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

from aidp.parsing import slide_render, slides  # noqa: E402

passed = failed = 0


def ok(name: str, condition: bool, extra: object = "") -> None:
    global passed, failed
    if condition:
        passed += 1
        print(f"  PASS {name}")
    else:
        failed += 1
        print(f"  FAIL {name} {extra}")


def picture_png() -> bytes:
    pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 400, 300), False)
    pixmap.clear_with(180)
    return pixmap.tobytes("png")


def build_deck() -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    title_only = prs.slide_layouts[5]
    bullets = prs.slide_layouts[1]

    # 1 — hidden, in front of everything, so a PDF without it shifts every page.
    hidden = prs.slides.add_slide(title_only)
    hidden.shapes.title.text = "Hidden slide"
    hidden._element.set("show", "0")  # noqa: SLF001

    # 2 — plain bullets.
    slide = prs.slides.add_slide(bullets)
    slide.shapes.title.text = "Key Principles"
    slide.placeholders[1].text_frame.text = "Cloud native\nSecure by design"

    # 3 — a drawn diagram: five boxes, four connectors.
    slide = prs.slides.add_slide(title_only)
    slide.shapes.title.text = "Proposed Logical Architecture"
    boxes = []
    for index, label in enumerate(["Client", "API Gateway", "Orders", "Payments", "Database"]):
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6 + index * 2.5), Inches(3), Inches(2), Inches(1)
        )
        box.text_frame.text = label
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
        boxes.append(box)
    for left, right in zip(boxes, boxes[1:]):
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, left.left + left.width, Inches(3.5), right.left, Inches(3.5)
        )
        connector.begin_connect(left, 3)
        connector.end_connect(right, 1)

    # 4 — text laid out in boxes, no arrows: text, already read.
    slide = prs.slides.add_slide(title_only)
    slide.shapes.title.text = "Assumptions"
    for index in range(6):
        box = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.6 + (index % 3) * 4), Inches(2 + (index // 3) * 2),
            Inches(3.5), Inches(1.5),
        )
        box.text_frame.text = f"Assumption {index + 1}"

    # 5 — a pasted diagram with a few lines drawn over it: the picture is the figure.
    slide = prs.slides.add_slide(title_only)
    slide.shapes.title.text = "Typical B2B Solution"
    slide.shapes.add_picture(io.BytesIO(picture_png()), Inches(1), Inches(1.5), Inches(11), Inches(5.5))
    for index in range(3):
        slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1 + index * 3), Inches(0.3),
                               Inches(1), Inches(0.5))
        slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1 + index * 3), Inches(0.9),
                                   Inches(2 + index * 3), Inches(0.9))

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


RAW = build_deck()
PRS = Presentation(io.BytesIO(RAW))

print("\nWhich slides are drawn diagrams")
bare = slides.read(RAW, render_diagrams=False)
picked = slides.diagram_slides(PRS, bare)
ok("only the boxes-and-arrows slide", sorted(picked) == [3], picked)
ok("its complexity counts the shapes that draw it", picked.get(3) == 9, picked.get(3))
ok("the pasted picture is still a figure of its own",
   [f.page for f in bare.figures if f.kind == "raster"] == [5], [f.page for f in bare.figures])
ok("with rendering off, no picture is taken and nothing is reported",
   not any(f.kind == "slide" for f in bare.figures) and not bare.unrendered_diagrams)

print("\nPhotographing it")
if slide_render.binary() is None:
    ok("LibreOffice is installed in this image", False, "soffice not found")
else:
    deck = slides.read(RAW)
    drawn = [f for f in deck.figures if f.kind == "slide"]
    ok("one slide figure", len(drawn) == 1, [(f.page, f.kind) for f in deck.figures])
    if drawn:
        figure = drawn[0]
        ok("of slide 3 — the hidden slide in front did not shift it", figure.page == 3, figure.page)
        ok("captioned with the slide's title", figure.caption == "Proposed Logical Architecture",
           figure.caption)
        ok("covering the whole slide", figure.bbox[0] == 0 and figure.bbox[2] > 900, figure.bbox)
        picture = fitz.Pixmap(figure.png)
        ok("a real PNG, at the diagram resolution", picture.width > 2000 and picture.height > 1100,
           (picture.width, picture.height))
        # The picture is of the diagram: its boxes are not the page's background.
        samples = {picture.pixel(x, int(picture.height * 0.47))
                   for x in range(0, picture.width, max(1, picture.width // 60))}
        ok("and it has the diagram drawn on it, not a blank page", len(samples) > 1, len(samples))
    ok("nothing reported missing", not deck.unrendered_diagrams, deck.render_error)
    ok("the pasted picture is still there beside it",
       any(f.kind == "raster" and f.page == 5 for f in deck.figures))

    pictures = slide_render.render(RAW, [2, 3], slide_count=len(PRS.slides))
    ok("any slide can be asked for, by number", sorted(pictures) == [2, 3], sorted(pictures))

print("\nWhen a picture cannot be taken")
real_binary = slide_render.binary
slide_render.binary = lambda: None
try:
    missing = slides.read(RAW)
finally:
    slide_render.binary = real_binary
ok("no LibreOffice: the deck is still read", len(missing.lines) > 0)
ok("and says which diagram has no picture, and why",
   missing.unrendered_diagrams == [3] and "not installed" in (missing.render_error or ""),
   (missing.unrendered_diagrams, missing.render_error))
ok("its other figures are kept", any(f.kind == "raster" for f in missing.figures))

real_run = subprocess.run


def stuck(*args, **kwargs):
    raise subprocess.TimeoutExpired(cmd="soffice", timeout=slide_render.TIMEOUT_SECONDS)


slide_render.binary = lambda: "/usr/bin/soffice"
subprocess.run = stuck
try:
    try:
        slide_render.render(RAW, [3], slide_count=5)
        ok("a stuck conversion is stopped", False, "no exception")
    except slide_render.RenderFailed as exc:
        ok("a stuck conversion is stopped, and says it took too long", "longer than" in str(exc),
           str(exc))
finally:
    subprocess.run = real_run
    slide_render.binary = real_binary

if slide_render.binary() is not None:
    try:
        slide_render.render(RAW, [3], slide_count=99)
        ok("a page count that is not the slide count is refused", False, "no exception")
    except slide_render.RenderFailed as exc:
        ok("a page count that is not the slide count is refused", "pages for 99 slides" in str(exc),
           str(exc))

ok("asking for nothing takes no picture", slide_render.render(RAW, [], slide_count=5) == {})

print(f"\n{passed} passed, {failed} failed")
sys.exit(1 if failed else 0)

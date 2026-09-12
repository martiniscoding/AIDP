"""Smoke test for rules read by a model — no database, no network.

The model is replaced by a scripted one, because what is under test is
everything around it — the part that has to hold however the model behaves:

  · source lines        paragraphs, bullets and table rows numbered once each,
                        a table's rows where it was written, its cells bound
                        to their columns, and no PDF table read twice
  · checking a reply    numbers that are not lines are dropped and counted;
                        a reply that points nowhere, covers too little or
                        does not parse is refused rather than half-believed
  · retrying            a refused range is retried, split when it is long,
                        and abandoned — not bisected forever — when it keeps
                        failing; a spent quota stops everything at once
  · slicing             every word of a rule comes from the source lines; part
                        labels and list glyphs are the only things removed
  · cross-checks        obligations nobody accounted for, obligations set
                        aside, readings that disagree, and rules only the old
                        parser found are all raised, at the right severity
  · the parse stage     adopts a reading that holds up, keeps the parser's
                        clauses whenever it does not, and says which happened

    PYTHONPATH=Workers python3 Workers/scripts/smoke_rules.py
"""

from __future__ import annotations

import io
import json
import os
import pathlib
import sys
import tempfile

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
os.environ.setdefault("STAGE", "parse")
os.environ.setdefault("DATABASE_URL", "postgresql://smoke/test")
# Pinned, not defaulted: the assertions below count readings, and a developer's
# Workers/.env is loaded into the same process.
os.environ["RULES_BY_MODEL"] = "on"
os.environ["RULES_READINGS"] = "2"

import fitz  # noqa: E402
from docx import Document  # noqa: E402

from aidp.ai import llm  # noqa: E402
from aidp.parsing import docs, furniture, rules, sections, spans, tables  # noqa: E402
from aidp.parsing.clauses import Clause  # noqa: E402
from aidp.stages import parse  # noqa: E402

passed = failed = 0


def ok(name: str, condition: bool, extra: object = "") -> None:
    global passed, failed
    if condition:
        passed += 1
        print(f"  PASS {name}")
    else:
        failed += 1
        print(f"  FAIL {name} {extra}")


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


def standard() -> bytes:
    d = Document()
    d.add_heading("Acme Security Standard", 0)
    d.add_paragraph("Document code SEC-01 · Version 2.0")
    d.add_heading("1. Purpose", 1)
    d.add_paragraph("This document defines the security standards for Acme systems.")
    d.add_heading("2. Access Control", 1)
    d.add_paragraph("Statement: All privileged access must use multi-factor authentication.")
    d.add_paragraph("Rationale: Stolen passwords are the most common route to a breach.")
    d.add_paragraph("Requirements:")
    d.add_paragraph("Administrators must enrol a hardware token", style="List Bullet")
    d.add_paragraph("Service accounts must not have interactive logins", style="List Bullet")
    d.add_heading("3. Password Policy", 1)
    d.add_paragraph("Passwords must meet the minimums below.")
    grid = d.add_table(rows=3, cols=2)
    for r, row in enumerate(
        [
            ["Control", "Minimum Standard"],
            ["Length", "12 characters"],
            ["Rotation", "Every 90 days"],
        ]
    ):
        for c, value in enumerate(row):
            grid.cell(r, c).text = value
    d.add_heading("4. Exceptions", 1)
    d.add_paragraph("Exceptions must be approved by the CISO and reviewed annually.")
    d.add_heading("5. Glossary", 1)
    d.add_paragraph("MFA: multi-factor authentication.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def at(source: list[rules.SourceLine], prefix: str) -> int:
    """The ordinal of the line starting with `prefix`, or of the ref `prefix`."""
    for line in source:
        if line.ref == prefix or line.text.startswith(prefix):
            return line.ordinal
    raise KeyError(prefix)


def good_reply(source: list[rules.SourceLine], *, exceptions: str | None = "process_rule") -> dict:
    n = lambda prefix: at(source, prefix)  # noqa: E731
    labels = [
        {"from": 0, "to": n("Document code"), "label": "furniture", "reason": "cover"},
        {"from": n("1. Purpose"), "to": n("This document"), "label": "scope", "reason": "purpose"},
        {"from": n("5. Glossary"), "to": n("MFA:"), "label": "reference", "reason": "glossary"},
    ]
    if exceptions:
        labels.append(
            {
                "from": n("4. Exceptions"),
                "to": n("Exceptions must"),
                "label": exceptions,
                "reason": "governs exceptions to the standard, not a design",
            }
        )
    return {
        "sections": [{"line": n(h), "depth": 1} for h in ("1. ", "2. ", "3. ", "4. ", "5. ")],
        "labels": labels,
        "rules": [
            {
                "heading": n("2. Access Control"),
                "from": n("2. Access Control"),
                "to": n("Service accounts"),
                "statement": [n("Statement:")],
                "rationale": [n("Rationale:")],
                "requirements": [
                    {"lines": [n("Administrators")], "strength": "must"},
                    {"lines": [n("Service accounts")], "strength": "must"},
                ],
                "guidance": [],
                "references": [],
            },
            {
                "heading": n("3. Password Policy"),
                "from": n("3. Password Policy"),
                "to": n("T1-r2"),
                "statement": [n("Passwords must")],
                "rationale": [],
                "requirements": [
                    {"lines": [n("T1-r1")], "strength": "must"},
                    {"lines": [n("T1-r2")], "strength": "should"},
                ],
                "guidance": [],
                "references": [n("MFA:")],
            },
        ],
    }


class Script:
    """A model that answers from a list, and remembers what it was asked."""

    def __init__(self, *replies) -> None:
        self.replies = list(replies)
        self.calls: list[tuple[int, int, int]] = []

    def __call__(self, *, document, title, first, last, whole, attempt) -> str:
        self.calls.append((first, last, attempt))
        reply = self.replies.pop(0) if len(self.replies) > 1 else self.replies[0]
        if isinstance(reply, Exception):
            raise reply
        if callable(reply):
            reply = reply(first, last)
        return reply if isinstance(reply, str) else json.dumps(reply)


def install(script: Script) -> Script:
    llm.read_rules = script
    llm.available = lambda: True
    parse._rules_schema = lambda: True
    return script


raw = standard()
read = docs.read(raw)
built, owner = sections.from_outline(read.lines, read.levels, document_title="Acme")
source = rules.word_source(read, owner)


# ---------------------------------------------------------------------------
print("\n1. Source lines from a Word standard")
refs = [line.ref for line in source]
ok("every line numbered once, in order",
   [ln.ordinal for ln in source] == list(range(len(source))))
ok("paragraph refs count paragraphs", refs[0] == "L1" and "L16" in refs, refs)
header = next((ln for ln in source if ln.ref == "T1-h"), None)
ok("table column line present", header is not None and header.text == "Control | Minimum Standard",
   header)
ok("table placed where it was written",
   at(source, "Passwords must") + 1 == at(source, "T1-h"), refs)
row = source[at(source, "T1-r1")]
ok("row cells bound to their columns",
   row.text == "Control: Length | Minimum Standard: 12 characters", row.text)
ok("row takes its name from its first cell", row.name == "Length", row.name)
bullet = source[at(source, "Administrators")]
ok("bullet kept as a bullet, glyph removed", bullet.kind == "bullet"
   and bullet.text == "Administrators must enrol a hardware token", bullet)
heading = source[at(source, "2. Access Control")]
ok("declared heading keeps its depth", heading.kind == "heading" and heading.depth == 1, heading)
password = next(i for i, s in enumerate(built) if s.title == "Password Policy")
ok("table rows fall in the section they were written under",
   row.section == password, (row.section, password))

rendered = rules.render(source, 0, len(source) - 1)
ok("render tags headings, bullets and rows",
   "[heading 1] 2. Access Control" in rendered
   and "[bullet] Administrators" in rendered
   and "[T1 row] Control: Length" in rendered, rendered[:400])


# ---------------------------------------------------------------------------
print("\n2. Checking a reply")
last = len(source) - 1
checked = rules.validate(
    {
        "sections": [{"line": 999, "depth": 1}],
        "labels": [{"from": 3, "to": 1, "label": "boilerplate", "reason": "legal text"}],
        "rules": [
            {"from": 4, "to": 9, "statement": ["5"], "rationale": [True], "requirements": [
                {"lines": [8, 400], "strength": "absolutely"}], "guidance": [], "references": []},
            {"from": 4, "to": 6, "statement": [15], "requirements": [], "guidance": []},
            {"from": 1, "to": 2, "statement": ["L7"], "requirements": []},
        ],
    },
    0,
    last,
)
ok("a number that is not a line is dropped and counted", checked.invalid >= 4, checked.invalid)
ok("string numbers are accepted", checked.rules and checked.rules[0].statement == [5],
   checked.rules[:1])
ok("unknown strength reads as must", checked.rules[0].requirements == [([8], "must")],
   checked.rules[0].requirements)
ok("a reversed range is put the right way round",
   (checked.labels[0].start, checked.labels[0].end) == (1, 3))
ok("an unknown label becomes 'other', keeping what it said",
   checked.labels[0].label == "other" and "boilerplate" in checked.labels[0].reason)
ok("a member far outside its span is not believed, and the empty rule goes",
   len(checked.rules) == 1, len(checked.rules))
ok("a reply this wrong is refused", rules.refusal(checked, 0, last) is not None)

thin = rules.validate(
    {"labels": [{"from": 0, "to": 2, "label": "furniture", "reason": "x"}]}, 0, last
)
ok("a reply covering too little is refused, saying how much",
   "accounted for only" in (rules.refusal(thin, 0, last) or ""), rules.refusal(thin, 0, last))
ok("a reply covering everything is believed",
   rules.refusal(rules.validate(good_reply(source), 0, last), 0, last) is None,
   rules.refusal(rules.validate(good_reply(source), 0, last), 0, last))


# ---------------------------------------------------------------------------
print("\n3. Reading, slicing and cross-reading")
script = install(Script(good_reply(source)))
outcome = rules.read(source, title="Acme", readings=2)
ok("first reading adopted", outcome.adopted is not None and outcome.adopted.number == 1)
ok("second reading made as the cross-check",
   len(outcome.attempts) == 2 and [c[2] for c in script.calls] == [1, 2], script.calls)

found = rules.to_clauses(source, outcome.reading, built)
clauses = [c for group in found.values() for c in group]
ok("one clause per rule", len(clauses) == 2, [c.title for c in clauses])
access = next((c for c in clauses if c.title == "Access Control"), None)
ok("title is the heading without its number", access is not None, [c.title for c in clauses])
if access:
    ok("statement sliced, label removed",
       access.statement == "All privileged access must use multi-factor authentication.",
       access.statement)
    ok("rationale sliced, label removed",
       access.rationale == "Stolen passwords are the most common route to a breach.",
       access.rationale)
    ok("each requirement its own item, verbatim", access.requirements == [
        "Administrators must enrol a hardware token",
        "Service accounts must not have interactive logins"], access.requirements)
    ok("marked as read by a model", access.origin == "model")
policy = next((c for c in clauses if c.title == "Password Policy"), None)
if policy:
    ok("table rows become requirements with their columns",
       policy.requirements[0] == "Control: Length | Minimum Standard: 12 characters",
       policy.requirements)
    ok("refs record where each part came from",
       policy.source["requirements"][1] == {"refs": ["T1-r2"], "strength": "should"}
       and policy.source["references"] == [source[at(source, "MFA:")].ref], policy.source)
    ok("clause lands in its own section", password in found and policy in found[password])

everything = " ".join(line.text for line in source)
words = [p for c in clauses for p in [c.statement, c.rationale, *c.requirements] if p]
ok("every word of every rule is in the document", all(p in everything for p in words),
   [p for p in words if p not in everything])


# ---------------------------------------------------------------------------
print("\n3b. A rule that is one row of a principle grid")
# A model can only point at a whole row. Pointed at for the statement, the
# rationale and the requirement at once, a principle-grid row once became three
# copies of itself; its cells are what hold the parts.
grid_doc = Document()
grid_doc.add_heading("Application Guidelines", 1)
grid = grid_doc.add_table(rows=2, cols=4)
for c, value in enumerate(["#", "Principle Name", "Statement and Rationale", "Implications"]):
    grid.cell(0, c).text = value
for c, value in enumerate([
    "1",
    "Reuse before Build",
    "Statement: Existing services must be evaluated before new ones are built. "
    "Rationale: Reuse lowers cost and risk.",
    "• Designs must record the reuse assessment",
]):
    grid.cell(1, c).text = value
buf = io.BytesIO()
grid_doc.save(buf)
grid_read = docs.read(buf.getvalue())
grid_sections, grid_owner = sections.from_outline(
    grid_read.lines, grid_read.levels, document_title="Guidelines"
)
grid_source = rules.word_source(grid_read, grid_owner)
row_at = at(grid_source, "T1-r1")
grid_reading = rules.validate(
    {
        "sections": [{"line": 0, "depth": 1}],
        "labels": [],
        "rules": [{
            "heading": None, "from": at(grid_source, "T1-h"), "to": row_at,
            "statement": [row_at], "rationale": [row_at],
            "requirements": [{"lines": [row_at], "strength": "must"}],
            "guidance": [], "references": [],
        }],
    },
    0,
    len(grid_source) - 1,
)
grid_clauses = [c for g in rules.to_clauses(grid_source, grid_reading, grid_sections).values()
                for c in g]
row_rule = grid_clauses[0] if grid_clauses else None
ok("one row, one rule", len(grid_clauses) == 1, grid_clauses)
if row_rule:
    ok("titled by the row's own name", row_rule.title == "Reuse before Build", row_rule.title)
    ok("statement taken from its cell, not the whole row",
       row_rule.statement == "Existing services must be evaluated before new ones are built.",
       row_rule.statement)
    ok("rationale taken from its cell", row_rule.rationale == "Reuse lowers cost and risk.",
       row_rule.rationale)
    ok("requirements taken from the implications cell",
       row_rule.requirements == ["Designs must record the reuse assessment"],
       row_rule.requirements)
    ok("still recorded as read by a model, with the row as its source",
       row_rule.origin == "model" and row_rule.source["statement"] == ["T1-r1"], row_rule.source)

from aidp.parsing import clauses as clause_parser  # noqa: E402

lined = clause_parser.from_table(grid_sections[0], [{
    "#": "2",
    "Principle Name": "Secure Data Capture",
    "Statement and Rationale": (
        "Statement\nAll data must be captured securely.\n\nRationale\nPrevents poor-quality data."
    ),
    "Implications": "Systems must validate input at entry\nSolutions must record capture metadata",
}])
ok("a Word cell's own lines divide the statement from the rationale",
   bool(lined) and lined[0].statement == "All data must be captured securely."
   and lined[0].rationale == "Prevents poor-quality data.", lined)
ok("implications listed one per line are separate requirements",
   bool(lined) and lined[0].requirements == [
       "Systems must validate input at entry", "Solutions must record capture metadata"], lined)
prose = clause_parser.from_table(grid_sections[0], [{
    "Principle Name": "Encryption",
    "Statement": "Data must be encrypted.",
    "Implications": "Keys must be rotated yearly. Keys must never be shared.",
}])
ok("a single run of implications is split into its sentences",
   bool(prose) and prose[0].requirements == [
       "Keys must be rotated yearly", "Keys must never be shared"], prose)

overlap = rules.validate(
    {"rules": [{
        "heading": at(source, "2. Access Control"), "from": at(source, "2. Access Control"),
        "to": at(source, "Service accounts"),
        "statement": [at(source, "Statement:")], "rationale": [at(source, "Statement:")],
        "requirements": [{"lines": [at(source, "Statement:")], "strength": "must"},
                         {"lines": [at(source, "Administrators")], "strength": "must"}],
        "guidance": [], "references": []}]},
    0,
    last,
)
repeated = [c for g in rules.to_clauses(source, overlap, built).values() for c in g][0]
ok("a line pointed at for two parts is used once, by the first",
   repeated.rationale == ""
   and repeated.requirements == ["Administrators must enrol a hardware token"],
   (repeated.rationale, repeated.requirements))


# ---------------------------------------------------------------------------
print("\n4. Retries, refusals and quota")
script = install(Script('{"sections": [', good_reply(source)))
outcome = rules.read(source, title="Acme", readings=1)
first = outcome.attempts[0]
ok("a cut-off reply is retried and the retry believed",
   outcome.adopted is first and [c.status for c in first.calls] == ["failed", "accepted"],
   [(c.status, c.error) for c in first.calls])
ok("one reading asked for, one made, when it succeeds", len(outcome.attempts) == 1)

script = install(Script("I cannot help with that"))
outcome = rules.read(source, title="Acme", readings=1)
ok("a reply that never parses is refused, not adopted", outcome.reading is None)
ok("a failed reading still gets one more try", len(outcome.attempts) == 2, len(outcome.attempts))
ok("each try retries its range once and stops", len(script.calls) == 4, script.calls)
ok("the reason is kept", bool(outcome.error), outcome.error)

script = install(Script(llm.QuotaExhausted("the model's per-day free-tier quota is exhausted")))
outcome = rules.read(source, title="Acme", readings=2)
ok("a spent quota stops at once", len(script.calls) == 1 and len(outcome.attempts) == 1,
   script.calls)
ok("and says so", "quota" in (outcome.error or ""), outcome.error)

long = [
    rules.SourceLine(ordinal=i, ref=f"L{i + 1}", kind="text", text=f"Line {i}", page=1, section=0)
    for i in range(250)
]


def cover(first: int, last: int) -> dict:
    return {"sections": [], "rules": [],
            "labels": [{"from": first, "to": last, "label": "other", "reason": "filler"}]}


script = install(Script('{"rules": [{"from": 0', cover))
outcome = rules.read(long, title="Long", readings=1)
ranges = [(c[0], c[1]) for c in script.calls]
ok("a long refused range is split in two and each half read", outcome.reading is not None
   and len(ranges) == 3 and ranges[0] == (0, 249) and ranges[1][0] == 0
   and ranges[2][1] == 249 and ranges[1][1] + 1 == ranges[2][0], ranges)
ok("the halves together account for every line",
   outcome.reading is not None and outcome.reading.covered() == set(range(250)))

script = install(Script("nope"))
outcome = rules.read(long, title="Long", readings=1)
ok("a range that keeps failing is abandoned within the call limit",
   outcome.reading is None
   and all(len(a.calls) <= rules.MAX_CALLS_PER_READING + 1 for a in outcome.attempts),
   [len(a.calls) for a in outcome.attempts])

saved = rules.MAX_LINES_PER_CALL
rules.MAX_LINES_PER_CALL = 60
headed = [
    rules.SourceLine(ordinal=i, ref=f"L{i + 1}", kind="heading" if i == 45 else "text",
                     text=f"Line {i}", page=1, section=0)
    for i in range(130)
]
pieces = rules.parts(headed)
rules.MAX_LINES_PER_CALL = saved
ok("a long document is divided at a heading, contiguously",
   pieces[0] == (0, 44) and pieces[-1][1] == 129
   and all(b[0] == a[1] + 1 for a, b in zip(pieces, pieces[1:], strict=False)), pieces)


# ---------------------------------------------------------------------------
print("\n5. Cross-checks raised as issues")


def findings(reply: dict, *, second: dict | None = None, parser: dict | None = None):
    install(Script(reply, second or reply))
    result = rules.read(source, title="Acme", readings=2)
    return {f.kind: f for f in rules.review(source, result, parser or {}, built)}


clean = findings(good_reply(source))
ok("a clean reading with an explained process rule raises only that",
   set(clean) == {"rules_obligation_set_aside"}
   and clean["rules_obligation_set_aside"].severity == "medium",
   {k: v.severity for k, v in clean.items()})

hidden = findings(good_reply(source, exceptions="furniture"))
ok("an obligation set aside as furniture is high",
   hidden.get("rules_obligation_set_aside") is not None
   and hidden["rules_obligation_set_aside"].severity == "high"
   and "CISO" in hidden["rules_obligation_set_aside"].detail, hidden)

unread = findings(good_reply(source, exceptions=None))
ok("an obligation nobody accounted for is high",
   unread.get("rules_lines_unread") is not None and unread["rules_lines_unread"].severity == "high"
   and "L14" in unread["rules_lines_unread"].detail, unread)

wider = good_reply(source, exceptions=None)
wider["rules"].append({
    "heading": at(source, "4. Exceptions"), "from": at(source, "4. Exceptions"),
    "to": at(source, "Exceptions must"), "statement": [at(source, "Exceptions must")],
    "rationale": [], "requirements": [], "guidance": [], "references": []})
differ = findings(good_reply(source), second=wider)
ok("readings that disagree about an obligation are raised",
   differ.get("rules_readings_differ") is not None
   and differ["rules_readings_differ"].severity == "high", differ)

install(Script(good_reply(source), llm.QuotaExhausted("per-day quota is exhausted")))
halfway = rules.read(source, title="Acme", readings=2)
missing = {f.kind: f for f in rules.review(source, halfway, {}, built)}
ok("a cross-check that never ran is said to have not run, not taken as agreement",
   halfway.adopted is not None and "rules_cross_check_missing" in missing
   and "quota" in missing["rules_cross_check_missing"].detail
   and "rules_readings_differ" not in missing, list(missing))

parser_view = {
    password: [Clause(ordinal=1, title="Password Policy",
                      statement="Passwords must meet the minimums below.")],
    next(i for i, s in enumerate(built) if s.title == "Exceptions"): [
        Clause(ordinal=1, title="Exceptions",
               statement="Exceptions must be approved by the CISO and reviewed annually.",
               page_start=1)],
}
parsed = findings(good_reply(source), parser=parser_view)
ok("a rule only the parser found is raised, with the model's reason",
   parsed.get("rules_parser_only") is not None
   and "governs exceptions" in parsed["rules_parser_only"].detail
   and parsed["rules_parser_only"].severity == "medium", parsed.get("rules_parser_only"))
ok("a rule both found is not", sum(1 for k in parsed if k == "rules_parser_only") == 1)


# ---------------------------------------------------------------------------
print("\n6. The parse stage")
reference = {"title": "Acme", "role": "reference"}

script = install(Script(good_reply(source)))
out = parse._parse_word(raw, reference, lambda: None)
kinds = {issue["kind"]: issue for issue in out.issues}
ok("a reading that holds up replaces the parser's clauses",
   out.clauses and all(c.origin == "model" for g in out.clauses.values() for c in g),
   {i: [c.origin for c in g] for i, g in out.clauses.items()})
ok("and puts them behind the confirmation gate", out.structure_inferred is True)
ok("and says what it did", "rules_by_model" in kinds, list(kinds))
ok("the parser's own guesses are not reported over it",
   "clause_not_extracted" not in kinds and "clause_inferred" not in kinds, list(kinds))
ok("source lines and readings kept for storage",
   out.source_lines is not None and len(out.source_lines) == len(source)
   and out.rule_outcome is not None and len(out.rule_outcome.attempts) == 2)
ok("nothing reported lost", "content_not_indexed" not in kinds, kinds.get("content_not_indexed"))

script = install(Script("not json at all"))
out = parse._parse_word(raw, reference, lambda: None)
kinds = {issue["kind"]: issue for issue in out.issues}
ok("a reading that fails keeps the parser's clauses",
   out.clauses and all(c.origin == "parser" for g in out.clauses.values() for c in g))
ok("and is reported high, and nothing is gated on it",
   kinds.get("rules_model_failed", {}).get("severity") == "high"
   and out.structure_inferred is False, list(kinds))

nothing = {"sections": [], "rules": [],
           "labels": [{"from": 0, "to": len(source) - 1, "label": "other", "reason": "?"}]}
install(Script(nothing))
out = parse._parse_word(raw, reference, lambda: None)
kinds = {issue["kind"] for issue in out.issues}
ok("zero rules against a parser that found some keeps the parser's",
   "rules_model_found_none" in kinds
   and all(c.origin == "parser" for g in out.clauses.values() for c in g), kinds)
ok("and the reading is recorded as not adopted",
   out.rule_outcome is not None and out.rule_outcome.adopted is None)

script = install(Script(good_reply(source)))
out = parse._parse_word(raw, {"title": "Design", "role": "assessed"}, lambda: None)
ok("a submitted design is never read for rules",
   not script.calls and out.source_lines is None
   and all(c.origin == "parser" for g in out.clauses.values() for c in g))

script = install(Script(good_reply(source)))
parse._rules_schema = lambda: False
out = parse._parse_word(raw, reference, lambda: None)
kinds = {issue["kind"] for issue in out.issues}
ok("an unmigrated database falls back to the parser and says how to fix it",
   not script.calls and "rules_schema_missing" in kinds and out.source_lines is None, kinds)


# ---------------------------------------------------------------------------
print("\n7. Source lines from a PDF")
from smoke_parse import build_pdf  # noqa: E402

path = pathlib.Path(tempfile.mkdtemp()) / "rules.pdf"
build_pdf(path)
pdf = fitz.open(path)
lines = spans.extract_lines(pdf)
profile = spans.profile_fonts(lines)
content, _ = furniture.strip(lines, pdf.page_count)
found_tables, regions = tables.extract_with_regions(pdf)
tree = sections.build(content, profile, document_title="Test Standards", skip_pages={2},
                      not_headings=regions)
pdf_lines = rules.pdf_source(
    content, tree, found_tables, regions,
    table_section=lambda t: parse._section_for_page(tree, t.page_start),
)
texts = [line.text for line in pdf_lines]
ok("numbered once, in order", [ln.ordinal for ln in pdf_lines] == list(range(len(pdf_lines))))
ok("a well-read table is given as its column line and rows",
   "Tier | Criticality | RPO | RTO" in texts
   and "Tier: Tier 1 | Criticality: Mission-critical | RPO: 0 to 2 hours | RTO: 1 to 2 hours"
   in texts, [t for t in texts if "Tier" in t])
ok("its cells are not read a second time as lines",
   "Mission-critical" not in texts and "Productivity" not in texts)
ok("an empty cell stays empty",
   "Tier: Tier 4 | Criticality: Function Specific" in texts, [t for t in texts if "Tier 4" in t])
tiers = next(i for i, s in enumerate(tree) if s.title == "Criticality Tiers")
ok("rows fall in the section they are printed under",
   all(ln.section == tiers for ln in pdf_lines if ln.kind.startswith("table")))
ok("headings are tagged as headings",
   any(ln.kind == "heading" and ln.text == "2. Governance Standards" for ln in pdf_lines),
   [ln.text for ln in pdf_lines if ln.kind == "heading"])
ok("the cover is kept, for the model to set aside as furniture",
   "Test Standards" in texts)
ok("running headers stay stripped", not any("Page" in t and "Test Standards" in t for t in texts))
pdf.close()

print(f"\n{passed} passed, {failed} failed")
sys.exit(1 if failed else 0)

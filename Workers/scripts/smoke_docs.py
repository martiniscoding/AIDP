"""Smoke test for the Word path — no database, no network.

Builds a .docx carrying the things that actually go wrong in one, then runs the
adapter over it and asserts on the result:

  · a house-renamed heading style          -> still a heading, at the right
                                              depth, because the outline level
                                              is inherited from the built-in
                                              style it is based on
  · a heading containing a colon            -> still a heading. The PDF path
                                              rejects those because typography
                                              cannot tell a heading from a
                                              "Label: sentence"; a declaration
                                              can, and second-guessing it would
                                              drop a real section
  · bulleted requirements                   -> each one its own obligation, not
                                              glued into a single run-on, which
                                              is what happens when the glyph is
                                              missing
  · a table written mid-document            -> lands in the section it was
                                              written under, not at the end
  · an empty cell in that table             -> stays empty rather than
                                              inheriting the value above it

    python3 Workers/scripts/smoke_docs.py
"""

from __future__ import annotations

import io
import os
import pathlib
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
# No database and no network: a real key in Workers/.env would otherwise send
# the reference standard below to a model. Rules read by a model have their
# own test, smoke_rules.py, with the model scripted.
os.environ["RULES_BY_MODEL"] = "off"

from docx import Document  # noqa: E402

from aidp.parsing import clauses as clause_parser  # noqa: E402
from aidp.parsing import docs, furniture, sections  # noqa: E402

passed = failed = 0


def ok(name: str, condition: bool, extra: object = "") -> None:
    global passed, failed
    if condition:
        passed += 1
        print(f"  PASS {name}")
    else:
        failed += 1
        print(f"  FAIL {name} {extra}")


def build() -> bytes:
    d = Document()
    # Where a real standards document states its classification, and where the
    # PDF it is exported to prints it on every page.
    d.sections[0].footer.paragraphs[0].text = "SENSITIVITY CLASSIFICATION: Internal Use"
    d.sections[0].header.paragraphs[0].text = "Security Standards | Page 1"
    d.add_heading("Security Standards", 0)

    # A house template renames its styles. Matching on the name would miss it;
    # the outline level is inherited from the style it is based on.
    house = d.styles.add_style("AIDP Heading 1", 1)
    house.base_style = d.styles["Heading 1"]
    d.add_paragraph("5. Access Control", style="AIDP Heading 1")

    d.add_heading("5.2 Key Management", level=2)
    d.add_paragraph("Statement: Cryptographic keys must be rotated every 90 days.")
    d.add_paragraph("Rationale: A long-lived key is a long-lived compromise.")
    d.add_paragraph("Requirements:")
    for text in (
        "Rotation is automated and requires no human action",
        "Rotation failures raise an alert within one hour",
        "Superseded keys are retained for 30 days then destroyed",
    ):
        d.add_paragraph(text, style="List Bullet")

    # A heading with a colon. The inferred path refuses these; a declared one
    # must not.
    d.add_heading("Scope: retail payments only", level=2)
    d.add_paragraph("This standard does not cover treasury systems.")

    d.add_heading("6. Recovery Objectives", level=1)
    table = d.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    for i, value in enumerate(["Tier", "RPO", "RTO"]):
        table.cell(0, i).text = value
    for r, row in enumerate(
        [
            ["Tier 2", "15 minutes", "1 hour"],
            ["Tier 3", "4 hours", "8 hours"],
            ["Tier 4", "", "24 hours"],
        ],
        start=1,
    ):
        for c, value in enumerate(row):
            table.cell(r, c).text = value

    buffer = io.BytesIO()
    d.save(buffer)
    return buffer.getvalue()


read = docs.read(build())

print("\n1. The document's own structure is taken at its word")
ok("title read off the Title style", read.title == "Security Standards", read.title)
levels = {read.lines[i].text: depth for i, depth in read.levels.items()}
ok("built-in heading found", levels.get("5.2 Key Management") == 2, levels)
ok("renamed house style still a heading", levels.get("5. Access Control") == 1, levels)
ok("a heading with a colon survives", levels.get("Scope: retail payments only") == 2, levels)
ok("body text is not promoted", "This standard does not cover treasury systems." not in levels)

print("\n2. Sections come out in order, with their paths")
built, owner = sections.from_outline(read.lines, read.levels, document_title=read.title or "")
titles = [s.title for s in built]
ok("four sections, the title not among them", len(built) == 4, titles)
ok("numbering lifted for display", built[0].number_text == "5", built[0].number_text)
ok("nesting is right", built[1].depth == 2 and built[1].heading_path.count("›") == 2,
   built[1].heading_path)

print("\n3. Bulleted requirements stay separate")
key = next(s for s in built if s.title == "Key Management")
found = clause_parser.from_prose(key)
ok("a clause was extracted", len(found) == 1, len(found))
if found:
    clause = found[0]
    ok("three requirements, not one", len(clause.requirements) == 3, clause.requirements)
    ok("statement captured", clause.statement.startswith("Cryptographic keys"), clause.statement)
    ok("rationale kept apart", clause.rationale.startswith("A long-lived key"), clause.rationale)

print("\n4. The table lands where it was written")
ok("one table", len(read.tables) == 1, len(read.tables))
if read.tables:
    anchor, table = read.tables[0]
    section_index = next(owner[i] for i in range(anchor, -1, -1) if owner[i] >= 0)
    ok("under Recovery Objectives", built[section_index].title == "Recovery Objectives",
       built[section_index].title)
    ok("headers bound", table.columns == ["Tier", "RPO", "RTO"], table.columns)
    ok("three data rows", len(table.rows) == 3, len(table.rows))
    tier4 = table.rows[-1]
    ok("Tier 4 is the last row", tier4["Tier"] == "Tier 4", tier4)
    ok("its empty RPO stays empty", tier4["RPO"] is None, tier4)
    ok("the value after it did not slide left", tier4["RTO"] == "24 hours", tier4)

print("\n5. The running footer is read, then kept out of the content")
ok("classification lifted", furniture.classification_in(
    [line.text for line in read.lines] + read.furniture) == "Internal Use")
body = " ".join(line.text for line in read.lines)
ok("footer text is not content", "SENSITIVITY CLASSIFICATION" not in body)
ok("header text is not content", "Page 1" not in body)

print("\n6. A document with no default paragraph style does not crash")
# Generators other than Word often omit the default paragraph style, and then
# every body paragraph reports its style as None. Three of four real client
# standards were built that way and crashed on their first body paragraph.
from docx.enum.style import WD_STYLE_TYPE  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

bare = Document()
for style in bare.styles:
    if style.element.get(qn("w:default")) == "1":
        style.element.attrib.pop(qn("w:default"))
ok("fixture really has no default paragraph style",
   bare.styles.default(WD_STYLE_TYPE.PARAGRAPH) is None)
bare.add_paragraph("Enterprise Information Security")
bare.add_heading("1. Introduction", level=1)
bare.add_paragraph("Statement: Systems must be assessed against these standards.")
bare.add_paragraph("Assessment happens before go-live", style="List Paragraph")
bare_buffer = io.BytesIO()
bare.save(bare_buffer)
try:
    bare_read = docs.read(bare_buffer.getvalue())
    ok("parsed without crashing", True)
    ok("the heading is still found",
       any(bare_read.lines[i].text == "1. Introduction" for i in bare_read.levels),
       bare_read.levels)
    ok("no heading is borrowed as the title", bare_read.title is None, bare_read.title)
except AttributeError as exc:
    ok("parsed without crashing", False, exc)


print("\n7. A cover before the first heading is not reported as lost text")
from aidp.stages import parse  # noqa: E402

WORD_ROW = {
    "id": "x",
    "title": "Security Standards",
    "role": "reference",
    "storageKey": "security.docx",
    "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def standard_with(opening: list[str]) -> bytes:
    built = Document()
    for text in opening:
        built.add_paragraph(text)
    built.add_heading("1. Access Control", level=1)
    for _ in range(4):
        built.add_paragraph(
            "Statement: Every privileged account must use multi-factor authentication."
        )
    out = io.BytesIO()
    built.save(out)
    return out.getvalue()


def lost_text(raw: bytes) -> list[dict]:
    result = parse._parse_word(raw, WORD_ROW, lambda: None)
    return [i for i in result.issues if i["kind"] == "content_not_indexed"]


cover = [
    "Enterprise Information Security",
    "Security Standards",
    "DOC-SEC-STD V-1.0",
    "Effective Date: [DD/MM/YYYY]",
]
cover_issues = lost_text(standard_with(cover))
ok("a cover is not counted as lost", not cover_issues, cover_issues)
intro = [
    (
        "This opening paragraph explains the purpose of the standard at length "
        "and is real content that belongs in the index. "
    )
    * 6
]
ok("real prose before the first heading still counts", bool(lost_text(standard_with(intro))))


print(f"\n{passed} passed, {failed} failed")
sys.exit(1 if failed else 0)
